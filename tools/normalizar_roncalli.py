"""Uso único/repetible: relee el Excel Roncalli y genera CSVs NORMALIZADOS
según las plantillas de `services/importar_plantillas.py` en importacion/normalizados/.

- RELACIÓN → estudiantes_relacion.csv (DNI provisional 90000001…, Nombres,
  Apellidos, BECA por color de fila, INICIO en ISO).
- BALANCE <MES> → tienda_balance_<mes>.csv (columnas plantilla Tienda +
  CATEGORIA vacía + FECHA = día 1 del mes de la hoja).
Los CSVs entran directo a Importar (Estudiantes/Tienda) sin errores de formato.
"""
import csv
import datetime
import glob
import os
import re
import unicodedata

MESES_HOJA = {"MAYO": 5, "JUNIO": 6, "JULIO": 7, "AGOSTO": 8,
              "SEPTIEMBRE": 9, "SEPTIEMRE": 9}


def norm(t):
    s = str(t or "").strip().upper()
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")
    return " ".join(s.split())


def num(txt):
    try:
        t = str(txt or "").strip().replace("s/", "").replace("S/", "").replace(",", "")
        return float(t) if t else 0.0
    except (ValueError, TypeError):
        return 0.0


def fill_beca(celda) -> str:
    try:
        f = celda.fill
        if not f or f.patternType in (None, "none"):
            return ""
        c = f.fgColor
        theme = str(getattr(c, "theme", "?")) if c.type != "rgb" else ""
        if theme == "6":
            return "1/2 BECA"
        if theme == "7":
            return "BECA COMPLETA"
        if c.type == "rgb":
            # Rellenos directos (verdes/celestes del XLSX real): misma
            # heurística que services/importar_historial._fill_verde_celeste.
            h = str(getattr(c, "rgb", "") or "")[-6:]
            r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
            if g >= 0x70 and g > r + 0x28 and g > b + 0x28:
                return "1/2 BECA"
            if b >= 0x70 and b > r + 0x28 and b > g:
                return "BECA COMPLETA"
    except Exception:
        pass
    return ""


def split_nombre(nombre):
    partes = str(nombre or "").strip().split()
    if not partes:
        return "", ""
    if len(partes) == 1:
        return partes[0], ""
    return partes[0], " ".join(partes[1:])


def fecha_inicio_depura(valor):
    import datetime as dt
    if isinstance(valor, dt.datetime):
        valor = valor.date()
    if isinstance(valor, dt.date):
        return "" if valor.year == 1900 else valor.isoformat()
    t = norm(valor)
    m = re.search(r"(\d{1,2})\s*/?\s*C\.?M\.?", t)
    if m:
        try:
            return dt.date(2026, 1, int(m.group(1))).isoformat()
        except ValueError:
            return ""
    m = re.search(r"INICIO\s+(\d{1,2})/(\d{1,2})", t)
    if m:
        try:
            return dt.date(2026, int(m.group(2)), int(m.group(1))).isoformat()
        except ValueError:
            return ""
    return ""


def main(origen=None, destino="importacion/normalizados"):
    import openpyxl
    if origen is None:
        cands = glob.glob("D:/Hp/Desktop/AcademiaFutbol/*.xlsx")
        if not cands:
            print("Sin xlsx origen")
            return
        origen = cands[0]
    os.makedirs(destino, exist_ok=True)
    wb = openpyxl.load_workbook(origen, read_only=True, data_only=True)

    # 1. RELACIÓN → estudiantes (se salta filas de título/leyenda/encabezado)
    ws = wb["RELACIÓN DE ALUMNOS"]
    filas = []
    dni = 90000001
    for r in ws.iter_rows(values_only=False):
        vals = [c.value for c in r]
        nombre = str(vals[1] or "").strip()
        if not nombre or norm(nombre).startswith(("NOMBRES", "ALUMNOS", "1/2", "BECA")):
            continue
        if norm(nombre) == "X CANCELADO":
            continue
        if vals[0] is None and not any(vals[2:11]):
            continue
        nombres, ape = split_nombre(nombre)
        beca = fill_beca(r[2])
        ini = ""
        for c in r[2:11]:
            v = c.value
            if isinstance(v, (datetime.datetime, datetime.date)) and getattr(v, "year", 0) != 1900:
                d = v.date() if isinstance(v, datetime.datetime) else v
                ini = d.isoformat()
                break
            f = fecha_inicio_depura(v)
            if f and not ini:
                ini = f
        filas.append([dni, nombres, ape, "", "", "", "", "", beca, ini])
        dni += 1
    with open(os.path.join(destino, "estudiantes_relacion.csv"), "w",
              encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["DNI", "Nombres", "Apellidos", "Fecha_Nacimiento", "Sexo",
                    "Telefono", "Correo", "Tipo_Documento", "BECA", "INICIO"])
        w.writerows(filas)
    print(f"estudiantes_relacion.csv: {len(filas)} filas")

    # 2. BALANCE → tienda (detecta fila de encabezado con PRODUCTOS)
    for hoja in [s for s in wb.sheetnames if norm(s).startswith("BALANCE")]:
        ws = wb[hoja]
        mes = 5
        for nombre_mes, numero in MESES_HOJA.items():
            if nombre_mes in norm(hoja):
                mes = numero
                break
        fecha = f"2026-{mes:02d}-01"
        header_idx = None
        headers = []
        todas = list(ws.iter_rows(values_only=True))
        for i, r in enumerate(todas):
            vals = [norm(v) for v in r]
            if any("PRODUCTOS" in v for v in vals):
                header_idx = i
                headers = vals
                break
        if header_idx is None:
            print(f"{hoja}: sin encabezado PRODUCTOS, omitida")
            continue

        def col(*nombres):
            for j, h in enumerate(headers):
                if h in nombres:
                    return j
            # MAYO trae "PRODUCTOS CANTIDAD" en una sola celda
            for j, h in enumerate(headers):
                for n in nombres:
                    if n and n in h and not any(
                            x in h for x in ("VENDIDO", "VENDIDAS", "VEND")
                            if "VEND" not in n):
                        return j
            return None

        i_nom = col("PRODUCTOS")
        i_cant = col("CANTIDAD")
        i_tot = col("COSTO TOTAL")
        i_uni = col("COSTO X UNIDAD")
        i_ven = col("COSTO VENTA")
        i_yape = col("YAPE", "P. YAPE", "P YAPE")
        i_efec = col("EFECTIVO", "P. EFECTIVO", "P EFECTIVO")
        i_vend = col("CANTIDAD VENDIDO", "CANTIDAD VENDIDAS", "UNIDADES VENDIDAS",
                     "CANTIDAD VENDIDAS", "VENDIDO")
        i_que = col("QUEDAN")
        if i_cant is None or i_cant == i_nom:
            # sin cabecera CANTIDAD (AGOSTO) o fusionada (MAYO): va al lado
            i_cant = i_nom + 1
        if i_tot is None and i_ven is not None:
            # AGOSTO sin cabeceras de costo: TOTAL y UNIT van antes de VENTA
            i_tot, i_uni = i_ven - 2, i_ven - 1
        out = []
        for r in todas[header_idx + 1:]:
            if i_nom is None or not str(r[i_nom] or "").strip():
                continue
            cant = int(num(r[i_cant]) if i_cant is not None else 0)
            tot = num(r[i_tot]) if i_tot is not None else 0
            uni = num(r[i_uni]) if i_uni is not None else (round(tot / cant, 2) if cant else 0)
            ven = num(r[i_ven]) if i_ven is not None else 0
            ya = num(r[i_yape]) if i_yape is not None else 0
            ef = num(r[i_efec]) if i_efec is not None else 0
            # JUNIO trae "15 UNIDADES": num() falla → 0; reintentar extrayendo dígitos
            if cant == 0 and i_cant is not None:
                m = re.search(r"[\d.]+", str(r[i_cant] or ""))
                cant = int(float(m.group())) if m else 0
            vd = 0
            if i_vend is not None:
                m = re.search(r"[\d.]+", str(r[i_vend] or ""))
                vd = int(float(m.group())) if m else 0
            qu = None
            if i_que is not None and str(r[i_que] or "").strip() not in ("", "None"):
                m = re.search(r"[\d.]+", str(r[i_que]))
                qu = int(float(m.group())) if m else None
            if qu is None:
                qu = cant - vd
            out.append([str(r[i_nom]).strip(), cant,
                        f"{tot:.2f}", f"{uni:.2f}", f"{ven:.2f}",
                        f"{ya:.2f}", f"{ef:.2f}", vd, qu, "", fecha])
        # fusiona duplicados del mismo mes (AGOSTO repite productos)
        unidas = {}
        orden = []
        for r in out:
            k = norm(r[0])
            if k in unidas:
                d = unidas[k]
                d[1] += r[1]  # CANTIDAD suma
                for j in (2, 5, 6):  # TOTAL, YAPE, EFECTIVO suman (montos)
                    d[j] = f"{float(d[j]) + float(r[j]):.2f}"
                d[7] += r[7]  # VENDIDO suma (unitario/venta: se conserva el primero)
            else:
                unidas[k] = [r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7],
                             None, r[9], r[10]]
                orden.append(k)
        out2 = []
        for k in orden:
            d = unidas[k]
            d[8] = d[1] - d[7]  # QUEDAN = CANTIDAD − VENDIDO
            out2.append(d)
        seguro = "".join(c if (c.isalnum() or c in ("_", "-")) else "_" for c in hoja)
        with open(os.path.join(destino, f"tienda_{seguro}.csv"), "w",
                  encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["PRODUCTOS", "CANTIDAD", "COSTO TOTAL", "COSTO X UNIDAD",
                        "COSTO VENTA", "YAPE", "EFECTIVO", "CANTIDAD VENDIDO",
                        "QUEDAN", "CATEGORIA", "FECHA"])
            w.writerows(out2)
        print(f"tienda_{seguro}.csv: {len(out2)} filas "
              f"(fusionadas {len(out) - len(out2)})")


if __name__ == "__main__":
    main()
