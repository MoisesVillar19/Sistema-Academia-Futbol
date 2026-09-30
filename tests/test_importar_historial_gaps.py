"""Historial gaps: beca rgb, split 150, completa autopagada, idempotencia."""
import datetime

from services import importar_historial as hist


def _xlsx(path, hojas, fills=None):
    """fills: {(hoja, fila_idx_1based, col_idx_1based): argb_hex}."""
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    primero = True
    for nombre, filas in hojas.items():
        ws = wb.active if primero else wb.create_sheet(nombre)
        ws.title = nombre
        for r in filas:
            ws.append(r)
        primero = False
    for (hoja, fr, fc), argb in (fills or {}).items():
        ws = wb[hoja]
        ws.cell(row=fr, column=fc).fill = PatternFill(
            start_color=argb, end_color=argb, fill_type="solid")
    wb.save(path)
    return path


REL = "RELACIÓN DE ALUMNOS"


def test_beca_verde_rgb_asigna_media(tmp_path):
    ruta = _xlsx(str(tmp_path / "h.xlsx"), {
        REL: [
            ["", "NOMBRES Y APELLIDOS", "2026 / MES"],
            ["", "", "MAYO"],
            [1, "VERDE UNO", "s/50 EF.", None, None, None, None, None, None, None,
             datetime.datetime(2026, 5, 2)],
        ],
    }, fills={(REL, 3, 3): "FF92D050"})
    imp = hist.ImportadorHistorial(id_usuario=1)
    res = imp.importar_relacion(ruta)
    assert not res["errores"], res["errores"]
    assert res["becas"] == 1, res
    from services import matricula_service
    from repositories import estudiante_repository
    est = estudiante_repository.obtener_todos()[0]
    mats = matricula_service.obtener_por_estudiante(est["id_estudiante"])
    becas = matricula_service.obtener_becas_por_matricula(mats[0]["id_matricula"])
    assert any("1/2" in b.get("beca_nombre", "") for b in becas)


def test_beca_celeste_rgb_autopaga(tmp_path):
    ruta = _xlsx(str(tmp_path / "h.xlsx"), {
        REL: [
            ["", "NOMBRES Y APELLIDOS", "2026 / MES"],
            [2, "CELESTE DOS", None, None, None, None, None, None, None, None,
             datetime.datetime(2026, 5, 19)],
        ],
    }, fills={(REL, 2, 3): "FF00B0F0"})
    imp = hist.ImportadorHistorial(id_usuario=1)
    res = imp.importar_relacion(ruta)
    assert not res["errores"], res["errores"]
    assert res["becas"] == 1, res
    # INICIO 05 → DIC: 8 cuotas autopagadas a tarifa
    assert res["cuotas"] == 8, res
    assert res["pagos"] == 8, res
    from repositories import pago_repository
    pagos = pago_repository.obtener_todos(limit=20)
    assert pagos and all("Beca completa" in (p.get("observacion") or "") for p in pagos)


def test_split_150_cuota_120_mas_uniforme_30(tmp_path):
    ruta = _xlsx(str(tmp_path / "h.xlsx"), {
        REL: [
            ["", "NOMBRES Y APELLIDOS", "2026 / MES"],
            ["", "", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SETIEMBRE"],
            [1, "SPLIT TRES", None, None, None, None, "S/150 Y.", None, None, None,
             datetime.datetime(2026, 5, 2)],
        ],
    })
    imp = hist.ImportadorHistorial(id_usuario=1)
    res = imp.importar_relacion(ruta)
    assert not res["errores"], res["errores"]
    from repositories import cuota_repository, venta_repository, producto_repository
    from database.connection import fetch_all
    cuotas = {c["periodo"]: dict(c) for c in fetch_all("SELECT * FROM cuota")}
    # SET adopta la auto-cuota (120 pagado); la cadena RN-014 genera OCT
    assert set(cuotas) == {"2026-09", "2026-10"}, cuotas
    assert cuotas["2026-09"]["monto_total"] == 120.0
    assert cuotas["2026-09"]["saldo"] == 0
    assert cuotas["2026-10"]["estado"] == "PENDIENTE"
    ventas = venta_repository.obtener_todos()
    assert len(ventas) == 1 and ventas[0]["monto_total"] == 30.0
    assert ventas[0]["tipo_venta"] == "TIENDA"
    assert producto_repository.obtener_por_codigo("UNIFORME-COM") is not None


def test_recorrida_no_duplica(tmp_path):
    hojas = {
        REL: [
            ["", "NOMBRES Y APELLIDOS", "2026 / MES"],
            ["", "", "MAYO"],
            [1, "REPITE CUATRO", "s/100 Y.", None, None, None, None, None, None, None,
             datetime.datetime(2026, 5, 2)],
        ],
    }
    ruta = _xlsx(str(tmp_path / "h.xlsx"), hojas)
    imp = hist.ImportadorHistorial(id_usuario=1)
    r1 = imp.importar_relacion(ruta)
    assert r1["estudiantes"] == 1
    imp2 = hist.ImportadorHistorial(id_usuario=1)
    r2 = imp2.importar_relacion(ruta)
    assert r2["estudiantes"] == 0 and r2["matriculas"] == 0, r2
    assert any("ya existe" in r or "sin cambios" in r for r in imp2.revision)
    from repositories import estudiante_repository
    assert len(estudiante_repository.obtener_todos()) == 1


def test_sinonimos_beca_csv():
    from services import importar_service
    id_media = importar_service._beca_id_por_nombre("media beca")
    id_completa = importar_service._beca_id_por_nombre("BECA ENTERA")
    assert id_media, "1/2 BECA del seed debe matchear 'media beca'"
    assert id_completa, "BECA COMPLETA del seed debe matchear 'BECA ENTERA'"
    assert importar_service._beca_id_por_nombre("No Existe XYZ") is None


def test_inicio_posterior_manda_primer_pago(tmp_path):
    """Caso Kenzo: INICIO octubre con pagos desde mayo → manda mayo."""
    import datetime
    ruta = _xlsx(str(tmp_path / "h.xlsx"), {
        REL: [
            ["", "NOMBRES Y APELLIDOS", "2026 / MES"],
            ["", "", "MAYO", "JUNIO", "JULIO", "AGOSTO"],
            [1, "KENZO CASO", "s/100 Y.", "s/100 Y.", "s/100 Y.", "s/100 Y.", None,
             None, None, None, datetime.datetime(2026, 10, 5)],
        ],
    })
    imp = hist.ImportadorHistorial(id_usuario=1)
    res = imp.importar_relacion(ruta)
    assert not res["errores"], res["errores"]
    assert any("manda 2026-05" in d for d in res["detalles"]), res["detalles"]
    from services import matricula_service
    from repositories import estudiante_repository
    est = estudiante_repository.obtener_todos()[0]
    mats = matricula_service.obtener_por_estudiante(est["id_estudiante"])
    assert mats and mats[0]["fecha_inicio"] == "2026-05-01", mats


def test_historial_adopta_existente(tmp_path):
    """Alumno del CSV previo: el historial completa meses sin duplicar."""
    import datetime
    from services import importar_service
    from repositories import estudiante_repository
    ok, _, r = importar_service.importar_estudiantes(
        [{"DNI": "44556677", "Nombres": "Adopta", "Apellidos": "Hist",
          "INICIO": "2026-05-02"}],
        id_usuario=1)
    assert ok and r["matriculas"] == 1
    ruta = _xlsx(str(tmp_path / "h.xlsx"), {
        REL: [
            ["", "NOMBRES Y APELLIDOS", "2026 / MES"],
            ["", "", "MAYO", "JUNIO"],
            [1, "ADOPTA HIST", "s/100 Y.", "s/100 EF.", None, None, None, None,
             None, None, datetime.datetime(2026, 5, 2)],
        ],
    })
    imp = hist.ImportadorHistorial(id_usuario=1)
    res = imp.importar_relacion(ruta)
    assert not res["errores"], res["errores"]
    assert res["estudiantes"] == 0, "no debe duplicar persona"
    assert res["matriculas"] == 0, "reutiliza la matrícula activa"
    assert res["pagos"] == 2, res
    assert len(estudiante_repository.obtener_todos()) == 1
    from repositories import cuota_repository
    per = {c["periodo"] for c in cuota_repository.obtener_todas_pendientes()
           } | {c["periodo"] for c in cuota_repository.obtener_por_matricula(1)}
    assert {"2026-05", "2026-06"} <= per, per


def test_auto_retiro_dos_vacios(tmp_path):
    import datetime
    ruta = _xlsx(str(tmp_path / "h.xlsx"), {
        REL: [
            ["", "NOMBRES Y APELLIDOS", "2026 / MES"],
            ["", "", "MAYO", "JUNIO"],
            [1, "VACIO RET", "s/100 Y.", None, None, None, None, None, None, None,
             datetime.datetime(2026, 5, 2)],
        ],
    })
    imp = hist.ImportadorHistorial(id_usuario=1)
    res = imp.importar_relacion(ruta)
    assert not res["errores"], res["errores"]
    from repositories import estudiante_repository
    est = estudiante_repository.obtener_todos()[0]
    assert est["estado"] == "RETIRADO", est
    assert est["fecha_retiro"] == "2026-05-31", est
    assert any("retiro automático" in d for d in res["detalles"])


def test_retiro_manual_con_fecha(crear_estudiante):
    from controllers import estudiante_controller
    from repositories import estudiante_repository
    eid = crear_estudiante()
    ok, _ = estudiante_controller.registrar_retiro(eid, "2026-06-24")
    assert ok
    assert estudiante_repository.obtener_por_id(eid)["fecha_retiro"] == "2026-06-24"


def test_historial_backfill_beca_en_recorrida(tmp_path):
    """Fila verde ya importada (sin beca): la re-corrida asigna la beca."""
    ruta = _xlsx(str(tmp_path / "h.xlsx"), {
        REL: [
            ["", "NOMBRES Y APELLIDOS", "2026 / MES"],
            ["", "", "MAYO", "JUNIO", "JULIO", "AGOSTO"],
            [1, "BACKFILL CINCO", None, None, None, "s/100 Y.", None, None, None, None,
             datetime.datetime(2026, 5, 2)],
        ],
    }, fills={(REL, 3, 6): "FF92D050"})
    from services import matricula_service
    from repositories import estudiante_repository
    imp = hist.ImportadorHistorial(id_usuario=1)
    r1 = imp.importar_relacion(ruta)
    assert r1["becas"] == 1, r1
    # quitar la beca para simular import sin beca, luego re-correr
    est = estudiante_repository.obtener_todos()[0]
    mats = matricula_service.obtener_por_estudiante(est["id_estudiante"])
    bid = r1_beca_id(mats[0]["id_matricula"])
    from repositories import matricula_beca_repository
    matricula_beca_repository.eliminar(mats[0]["id_matricula"], bid)
    imp2 = hist.ImportadorHistorial(id_usuario=1)
    r2 = imp2.importar_relacion(ruta)
    assert r2["estudiantes"] == 0, r2
    assert r2["becas"] == 1, r2
    assert any("asignada a matrícula activa" in d for d in r2["detalles"]), r2


def r1_beca_id(id_matricula):
    from services import matricula_service
    becas = matricula_service.obtener_becas_por_matricula(id_matricula)
    assert becas
    return becas[0]["id_beca"]
