"""UI historial: tipo visible, sin mapeo, resultados con revisión."""
import pytest

pytestmark = pytest.mark.ui


def _textos(w):
    out = []
    vistos = set()

    def _rec(x):
        try:
            key = str(x)
        except Exception:
            return
        if key in vistos:
            return
        vistos.add(key)
        try:
            t = x.cget("text")
            if t:
                out.append(str(t))
        except Exception:
            pass
        try:
            for h in x.winfo_children():
                _rec(h)
        except Exception:
            pass

    _rec(w)
    return " ".join(out)


def _xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RELACIÓN DE ALUMNOS"
    ws.append(["", "NOMBRES Y APELLIDOS", "2026 / MES"])
    ws.append(["", "", "MAYO"])
    ws.append([1, "UI HIST QA", "s/100 Y.", None, None, None, None, None, None, None,
               __import__("datetime").datetime(2026, 5, 4)])
    ws2 = wb.create_sheet("INGRESOS")
    ws3 = wb.create_sheet("VENTA UNIFORME")
    wb.save(path)
    return path


def test_tipo_historial_sin_mapeo(crear_vista, usuario_admin):
    from views.importar.importar_view import ImportarView
    vista = crear_vista(ImportarView)
    vista.seg_tipo.set("Historial")
    vista._on_tipo_cambiar("Historial")
    assert vista._es_historial()
    vista._configurar_mapeo()
    vista.update_idletasks()
    assert "No requiere mapeo" in _textos(vista.scroll_mapeo)


def test_ejecutar_historial_notifica(crear_vista, usuario_admin, tmp_path, monkeypatch):
    import tkinter.messagebox as mb
    from views.importar.importar_view import ImportarView
    ruta = _xlsx(str(tmp_path / "h.xlsx"))
    vista = crear_vista(ImportarView)
    vista._archivo_actual = ruta
    vista._tipo_importacion = "Historial"
    monkeypatch.setattr(mb, "askyesno", lambda *a, **k: True)
    monkeypatch.setattr(mb, "showwarning", lambda *a, **k: None)
    monkeypatch.setattr(mb, "showerror", lambda *a, **k: None)
    # _avisar_resultado usa showinfo/showwarning según el resultado; sin
    # mockear el modal real cuelga la suite (cuelgue histórico de este archivo).
    monkeypatch.setattr(mb, "showinfo", lambda *a, **k: None)
    vista._ejecutar_historial()
    vista.update_idletasks()
    todo = _textos(vista.scroll_resultados)
    assert "Estudiantes creados" in todo
    assert "Matrículas creadas" in todo
    assert "Pagos registrados" in todo


def test_historial_rechaza_no_xlsx(crear_vista, usuario_admin, tmp_path, monkeypatch):
    import tkinter.messagebox as mb
    from views.importar.importar_view import ImportarView
    vista = crear_vista(ImportarView)
    vista._archivo_actual = str(tmp_path / "nada.csv")
    avisos = []
    monkeypatch.setattr(mb, "showwarning", lambda *a, **k: avisos.append(1))
    vista._ejecutar_historial()
    assert avisos


def test_secretaria_sin_importar(crear_vista, usuario_secretaria):
    from views.importar.importar_view import ImportarView
    vista = crear_vista(ImportarView)
    assert "denegado" in _textos(vista).lower()
