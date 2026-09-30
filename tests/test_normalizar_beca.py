"""Normalizador: beca por color directo rgb (JAHZIEL quedaba vacío)."""
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))


def _celda(argb):
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"].fill = PatternFill(start_color=argb, end_color=argb,
                                fill_type="solid")
    return ws["A1"]


def test_fill_beca_rgb():
    from tools.normalizar_roncalli import fill_beca
    assert fill_beca(_celda("FF92D050")) == "1/2 BECA"
    assert fill_beca(_celda("FF00B0F0")) == "BECA COMPLETA"
    assert fill_beca(_celda("FFFFFFFF")) == ""
